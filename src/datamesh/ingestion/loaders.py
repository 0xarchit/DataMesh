import csv
import glob
import os
import re

from openpyxl import load_workbook

from datamesh.core import llm
from datamesh.core.schema import Record, parse_amount


def load_csv(path):
    records = []
    with open(path, newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            records.append(
                Record(
                    source_type="csv",
                    source_file=os.path.basename(path),
                    vendor=row.get("company_name", "").strip(),
                    city=row.get("registered_city", "").strip(),
                    tax_id=row.get("gstin", "").strip(),
                    confidence=1.0,
                )
            )
    return records


def load_excel(path):
    records = []
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    for raw in rows[1:]:
        row = dict(zip(header, raw))
        records.append(
            Record(
                source_type="xlsx",
                source_file=os.path.basename(path),
                invoice_id=str(row.get("invoice_id", "")).strip(),
                vendor=str(row.get("vendor", "")).strip(),
                amount=parse_amount(row.get("amount_inr")),
                date=str(row.get("date", "")).strip(),
                tax_id=str(row.get("gstin", "")).strip(),
                confidence=1.0,
            )
        )
    return records


def _field(text, label):
    m = re.search(re.escape(label) + r"\s*:?\s*(.+)", text)
    return m.group(1).strip() if m else ""


def load_pdf(path):
    from pypdf import PdfReader

    text = "\n".join((page.extract_text() or "") for page in PdfReader(path).pages)
    return [
        Record(
            source_type="pdf",
            source_file=os.path.basename(path),
            invoice_id=_field(text, "Invoice No"),
            vendor=_field(text, "Supplier"),
            amount=parse_amount(_field(text, "Amount Payable")),
            date=_field(text, "Date"),
            city=_field(text, "Place of Supply"),
            tax_id=_field(text, "GSTIN"),
            confidence=0.95,
        )
    ]


def load_image(path, on_token=None):
    data = llm.vision_extract(path, on_token=on_token)
    return [
        Record(
            source_type="receipt",
            source_file=os.path.basename(path),
            invoice_id=str(data.get("receipt_no", "")).strip(),
            vendor=str(data.get("vendor", "")).strip(),
            amount=parse_amount(data.get("amount")),
            date=str(data.get("date", "")).strip(),
            confidence=0.8,
        )
    ]


def load_all(raw_dir, use_llm=True, emit=None, on_token=None):
    say = emit or (lambda s: None)
    records = []

    part = load_csv(os.path.join(raw_dir, "customers.csv"))
    say("csv customers.csv -> {} vendor rows".format(len(part)))
    records += part

    part = load_excel(os.path.join(raw_dir, "transactions.xlsx"))
    say("xlsx transactions.xlsx -> {} ledger rows".format(len(part)))
    records += part

    pdfs = sorted(glob.glob(os.path.join(raw_dir, "invoices", "*.pdf")))
    say("pdf invoices -> {} files".format(len(pdfs)))
    for i, path in enumerate(pdfs, start=1):
        part = load_pdf(path)
        if i % 50 == 0 or i == len(pdfs):
            say("pdf {}/{} parsed".format(i, len(pdfs)))
        records += part

    images = sorted(glob.glob(os.path.join(raw_dir, "receipts", "*.png")))
    if images:
        if use_llm and llm.available():
            say("vision receipts -> {} files, local model live".format(len(images)))
            for i, path in enumerate(images, start=1):
                name = os.path.basename(path)
                say("LLM vision {}/{} {}".format(i, len(images), name))
                try:
                    part = load_image(path, on_token=on_token)
                    say("vision {} ok: {} INR {} ({})".format(name, part[0].vendor, part[0].amount, part[0].invoice_id))
                    records += part
                except Exception as exc:
                    say("receipt skipped {}: {}".format(name, exc))
        else:
            say("vision receipts -> {} files skipped, LLM off".format(len(images)))
    say("ingest done -> {} records".format(len(records)))
    return records
